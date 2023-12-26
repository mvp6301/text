from flask import Flask, render_template, request, redirect, url_for,send_file  
from googletrans import Translator
from gtts import gTTS
import os
import cv2
import pytesseract
from PyPDF2 import PdfReader
from docx import Document
from sumy.parsers.plaintext import PlaintextParser
from sumy.nlp.tokenizers import Tokenizer
from sumy.summarizers.lsa import LsaSummarizer
import base64
import nltk
import io
from io import BytesIO
import textwrap
from fpdf import FPDF


nltk.download('punkt')

app = Flask(__name__)

LANGUAGE_MAPPING = {
    'telugu': 'te',
    'english': 'en',
    'assamese': 'as',
    'bengali': 'bn',
    'bodo': 'brx',
    'dogri': 'doi',
    'gujarati': 'gu',
    'hindi': 'hi',
    'kannada': 'kn',
    'kashmiri': 'ks',
    'konkani': 'kok',
    'maithili': 'mai',
    'malayalam': 'ml',
    'manipuri': 'mni',
    'marathi': 'mr',
    'nepali': 'ne',
    'odia': 'or-IN',
    'punjabi': 'pa',
    'sanskrit': 'sa',
    'santali': 'sat',
    'sindhi': 'sd',
    'tamil': 'ta',
    'telugu': 'te',
    'urdu': 'ur',
}

summarized_text=" "
def extract_text_from_image(scanned_image):
    gray_image = cv2.cvtColor(scanned_image, cv2.COLOR_RGB2GRAY)
    text = pytesseract.image_to_string(gray_image)
    print("Extracted Text from Image:")
    print(text)
    return text


def extract_text_from_pdf(pdf_path):
    with open(pdf_path, 'rb') as file:
        pdf_reader = PdfReader(file)
        text = ""
        for page_num in range(len(pdf_reader.pages)):
            page = pdf_reader.pages[page_num]
            text += page.extract_text()
        print("Extracted Text from PDF:")
        print(text)
        return text


def extract_text_from_text_file(text_file_path):
    with open(text_file_path, 'r') as file:
        text = file.read()
        print("Text from Text File:")
        print(text)
        return text


def extract_text_from_word_document(docx_path):
    doc = Document(docx_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + '\n'
    print("Text from Word Document:")
    print(text)
    return text


def summarize_text_sumy(text, sentences_count=10):
    parser = PlaintextParser.from_string(text, Tokenizer("english"))
    summarizer = LsaSummarizer()
    summary = summarizer(parser.document, sentences_count=sentences_count)
    summary_text = " ".join(str(sentence) for sentence in summary)
    print("Summary:")
    print(summary_text)
    return summary_text


def translate_text(text, target_language='en'):
    translator = Translator()

    # Handle empty text
    if not text:
        print("Error: Empty text provided for translation.")
        return ""

    try:
        translation = translator.translate(text, dest=target_language)
        translated_text = translation.text
        print(translated_text)
        return translated_text
    except Exception as e:
        print(f"Error during translation: {e}")
        return ""


def text_to_speech(text, language='en'):
    tts = gTTS(text=text, lang=language, slow=False)
    audio_buffer = BytesIO()
    tts.write_to_fp(audio_buffer)
    audio_data = audio_buffer.getvalue()
    return audio_data

def text_to_pdf(text, filename="summarized_text.pdf"):
    a4_width_mm = 210
    pt_to_mm = 0.35
    fontsize_pt = 10
    fontsize_mm = fontsize_pt * pt_to_mm
    margin_bottom_mm = 10
    character_width_mm = 7 * pt_to_mm
    width_text = a4_width_mm / character_width_mm

    pdf = FPDF(orientation='P', unit='mm', format='A4')
    pdf.set_auto_page_break(True, margin=margin_bottom_mm)
    pdf.add_page()
    pdf.set_font(family='Courier', size=fontsize_pt)
    splitted = text.split('\n')

    for line in splitted:
        lines = textwrap.wrap(line, width_text)

        if len(lines) == 0:
            pdf.ln()

        for wrap in lines:
            pdf.cell(0, fontsize_mm, wrap.encode('latin-1', 'replace').decode('latin-1'), ln=1)

    pdf.output(filename, 'F')
    return filename  # Return the suggested filename


@app.route('/')
def index():
    return render_template('index2.html')


@app.route('/process', methods=['POST'])
def process():
    user_choice = request.form.get('user_choice')
    if user_choice == 'camera':
        return redirect(url_for('camera'))
    elif user_choice == 'file':
        return redirect(url_for('file'))
    else:
        return "Invalid source option. Please choose 'camera' or 'file'."


@app.route('/camera')
def camera():
    return render_template('camera2.html', language_mapping=LANGUAGE_MAPPING)


@app.route('/camera_processing', methods=['POST'])
def camera_processing():
    camera_index = int(request.form.get('camera_index', 0))
    sentences_count = int(request.form.get('sentences_count', 3))
    target_language = request.form.get('target_language')
    translation_choice = request.form.get('translation_choice')

    cap = cv2.VideoCapture(camera_index)

    if not cap.isOpened():
        print("ERROR: Could not open the camera")
        return "Error: Could not open the camera"

    captured_images = []  # List to store captured images

    capturing = True
    while capturing:
        ret, frame = cap.read()
        cv2.imshow("Camera", frame)

        key = cv2.waitKey(1) & 0xFF

        if key == ord('c'):  # Press 'c' to capture an image
            cv2.imwrite("scanned_image.png", frame)
            print("Image captured successfully as 'scanned_image.png'.")
            scanned_image = cv2.imread("scanned_image.png")
            scanned_image = cv2.cvtColor(scanned_image, cv2.COLOR_BGR2RGB)  # Convert BGR to RGB
            captured_images.append(scanned_image)

        elif key == ord('s'):  # Press 's' to stop capturing
            capturing = False

    cap.release()
    cv2.destroyAllWindows()

    if captured_images:  # Check if any images were captured
        # Extract text from all captured images
        extracted_texts = [extract_text_from_image(image) for image in captured_images]
        combined_text = '\n'.join(extracted_texts)

        # Process the combined text
        summarized_text = summarize_text_sumy(combined_text, sentences_count)
        translated_text = translate_text(summarized_text, target_language) if translation_choice == 'yes' else summarized_text

        if not translated_text:  # Check if translated text is empty
            print("Error: Empty text provided for translation.")
            return "Error: Empty text provided for translation."

        # Generate audio data
        audio_data = text_to_speech(translated_text, language=target_language)
        audio_data_base64 = base64.b64encode(audio_data).decode('utf-8')

        return render_template('result1.html', extracted_text=combined_text, summarized_text=summarized_text,
                               translated_text=translated_text, audio_data=audio_data_base64,
                               language_mapping=LANGUAGE_MAPPING)
    else:
        return "Error: No images captured."




@app.route('/file')
def file():
    return render_template('file.html', language_mapping=LANGUAGE_MAPPING)


@app.route('/process_file', methods=['POST'])
def process_file():
    files = request.files.getlist('file')  # Use getlist to handle multiple file uploads
    sentences_count = int(request.form['sentences_count'])
    target_language = request.form['target_language']
    translation_choice = request.form['translation_choice']
    translated_text = ""
    if files:
        extracted_texts = []

        for file in files:
            # Save each file temporarily
            file_path = "temp_file" + os.path.splitext(file.filename)[-1]
            file.save(file_path)

            file_extension = os.path.splitext(file_path)[-1].lower()
            if file_extension == '.pdf':
                text = extract_text_from_pdf(file_path)
            elif file_extension == '.txt':
                text = extract_text_from_text_file(file_path)
            elif file_extension == '.docx':
                text = extract_text_from_word_document(file_path)
            elif file_extension in ['.png', '.jpg', '.jpeg', '.bmp']:
                scanned_image = cv2.imread(file_path)
                scanned_image = cv2.cvtColor(scanned_image, cv2.COLOR_BGR2RGB)
                text = extract_text_from_image(scanned_image)
            else:
                print(f"Unsupported file format: {file_extension} for file {file_path}")
                continue  # Skip to the next iteration

            os.remove(file_path)  # Remove the temporary file
            extracted_texts.append(text)

        all_text = '\n'.join(extracted_texts)

        summarized_text = summarize_text_sumy(all_text, sentences_count)
        pdf_filename = "summarized_text.pdf"
        #pdf_path = os.path.join(app.root_path, "static", pdf_filename)
        suggested_pdf_filename = text_to_pdf(summarized_text, filename=pdf_filename)
        if translation_choice == 'yes': 
            translated_text = translate_text(summarized_text, target_language) 
            audio_data = text_to_speech(translated_text, language=target_language)
        else:
            #translated_text = all_text  # Assign a value to translated_text in the else block
            target_language='en'
            audio_data = text_to_speech(all_text, language=target_language)

        # Generate audio data
        audio_data_base64 = base64.b64encode(audio_data).decode('utf-8')

        return render_template('result1.html', extracted_text=all_text, summarized_text=summarized_text,
                            translated_text=translated_text, audio_data=audio_data_base64,
                            language_mapping=LANGUAGE_MAPPING, pdf_filename= suggested_pdf_filename)
    else:
        return "Error: No files provided."
    
@app.route('/download_pdf/<filename>')
def download_pdf(filename):
    # Decode the filename
    #decoded_filename = unquote(filename)

    # Construct the full path to the secure location
    #pdf_path = os.path.join(app.root_path, "secure_folder", decoded_filename)
    pdf_path='summarized_text.pdf'

    # Check if the file exists
    if not os.path.exists(pdf_path):
        return "File not found"

    response = send_file(pdf_path, as_attachment=True)
    print(response.headers)
    return response

    
if __name__ == "__main__":
    app.run(debug=True)
